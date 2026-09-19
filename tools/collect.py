"""Local, resumable collection. Never reads or uploads account exports."""
import argparse, base64, collections, hashlib, io, json, re, time, urllib.request, urllib.error
from pathlib import Path
from datetime import datetime, timezone

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'data'
SNAP = DATA / 'source-snapshots'
CDN = 'https://yys.res.netease.com/pc/zt/20161108171335/data'
CATALOG = 'https://g37simulator.webapp.163.com/get_heroid_list'
STATIC = 'https://yys.res.netease.com/pc/zt/20161108171335/js/app/all_shishen.json'
API = 'https://api.fireschain.org/onmyoji/v1/team-code/decode'
ALLOWED = {'g37simulator.webapp.163.com','yys.res.netease.com','yys.163.com','api.fireschain.org','raw.githubusercontent.com','api.github.com','api.bilibili.com'}
last_request = 0

def now(): return datetime.now(timezone.utc).isoformat()
def save(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix+'.tmp')
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    tmp.replace(path)
def read(path, default=None):
    return json.loads(path.read_text(encoding='utf-8')) if path.exists() else default

class NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, *args, **kwargs): return None

def request(url, label, body=None, max_bytes=12*1024*1024):
    global last_request
    from urllib.parse import urlparse
    parsed = urlparse(url)
    if parsed.scheme != 'https' or parsed.hostname not in ALLOWED or parsed.username:
        raise ValueError('URL outside allowlist')
    evidence = SNAP / (label+'.json')
    if evidence.exists():
        m=read(evidence)
        p=SNAP/(label+'.bin')
        return m, p.read_bytes() if p.exists() else b''
    time.sleep(max(0, 1-(time.monotonic()-last_request)))
    last_request=time.monotonic()
    m={'url':url,'retrievedAt':now(),'httpStatus':None,'status':'network_error'}
    raw=b''
    headers={'User-Agent':'OnmyojiLocalLibrary/1.0','Accept':'application/json,image/*,text/html;q=0.9'}
    if body is not None: headers['Content-Type']='application/json'
    req=urllib.request.Request(url, data=json.dumps(body).encode() if body is not None else None, headers=headers)
    try:
        with urllib.request.build_opener(NoRedirect).open(req, timeout=22) as r:
            raw=r.read(max_bytes+1)
            m.update(httpStatus=r.status, contentType=r.headers.get('Content-Type'), cors=r.headers.get('Access-Control-Allow-Origin'))
            if len(raw)>max_bytes: raise ValueError('response-too-large')
            m['status']='received'
    except urllib.error.HTTPError as e:
        m.update(httpStatus=e.code,status={403:'access_denied',429:'rate_limited',404:'missing_resource'}.get(e.code,'http_error'),retryAfter=e.headers.get('Retry-After'))
        raw=e.read(8192)
    except Exception as e: m['error']=str(e)[:300]
    if raw:
        m.update(sha256=hashlib.sha256(raw).hexdigest(),bytes=len(raw))
        SNAP.mkdir(parents=True,exist_ok=True)
        (SNAP/(label+'.bin')).write_bytes(raw)
    save(evidence,m)
    return m,raw

def jsonp(raw):
    s=raw.decode('utf-8-sig').strip()
    if s.startswith(('cb(', 'callback(')):
        s=s[s.index('(')+1:]
        s=re.sub(r'\)\s*;?\s*$','',s)
    return json.loads(s)

def extract_excel():
    import openpyxl
    def route(category, subcategory, dungeon=None):
        return {'category':category,'subcategory':subcategory,'dungeon':dungeon or subcategory}
    def dashen_paths(group, title):
        if group=='近期更新': return []  # 更新引用复用正文的用途，不新造一个副本。
        actor=next((x for x in ['藤原道长','源赖光'] if x in group),None)
        if actor: return [route('周常','英杰试炼',actor+' · '+('经验本' if '经验' in title else 'PVP技能本' if 'PVP' in title else 'PVE技能本' if 'PVE' in title else '技能本'))]
        if '突破' in group: return [route('日常','结界突破')]
        if '契灵' in title:
            if '首领' in title: return [route('周常','契灵首领')]
            spirit=next((x for x in ['镇墓兽','月魔兔','薙魂','针女','狐火','火灵','茨球','小黑'] if x in title),'探查')
            return [route('契灵',spirit)]
        if '极逢魔' in title: return [route('极逢魔','荒骷髅' if '周四' in title and '除了' not in title else '通用')]
        if '彼世逢魔' in title: return [route('周常','彼世逢魔')]
        if '逢魔' in title: return [route('普通逢魔','鬼灵歌伎' if re.search('歌[姬伎]',title) else '通用')]
        if '寮' in group:
            if re.search('提赏金|双拉|清杂|有雀',title): return [route('道馆','道馆进攻')]
            if all(x in title for x in ['麒麟','狭间','退治']): return [route('寮活动',x) for x in ['麒麟','狭间暗域','首领退治']]
            for key,sub in [('狭间','狭间暗域'),('麒麟','麒麟'),('退治','首领退治')]:
                if key in title: return [route('寮活动',sub)]
            return [route('寮活动','僵尸寮通用')]
        for pattern,category,sub in [
            ('魂海.?P1','御魂','魂海P1'),('魂海.?P2','御魂','魂海P2'),('魂十','御魂','魂十'),('魂土','御魂','魂土'),('魂王','御魂','魂王'),('虚无','御魂','魂主'),('业原火','御魂','业原火'),('御灵','御魂','御灵'),
            ('阴界','寮活动','阴界之门'),('真蛇','周常','真蛇'),('秘闻|百战','周常','秘闻'),('六道','周常','六道'),
            ('地域鬼王','日常','地鬼'),('师徒','日常','师徒副本'),('金币','日常','金币妖怪'),('经验','日常','经验妖怪'),('探索|悬赏','日常','探索 / 悬赏')]:
            if re.search(pattern,title): return [route(category,sub)]
        return [route(group.removesuffix('副本'),title)]
    entries=[]; sources=[]; sheets=[]
    names=['大肾石长姬版本PVE阵容码大全 最新.xlsx','01 阴阳师阵容码合集  洛天依版本260902.xlsx']
    for name in names:
        p=ROOT/name; w=openpyxl.load_workbook(p,data_only=False); luotianyi=name.startswith('01 ')
        source={'source':name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'sheets':[]}
        start=len(entries)
        for s in w:
            merged={}
            for region in s.merged_cells.ranges:
                value=s.cell(region.min_row,region.min_col).value
                for row in range(region.min_row,region.max_row+1):
                    for col in range(region.min_col,region.max_col+1): merged[row,col]=value
            def value(row,col):
                v=merged.get((row,col),s.cell(row,col).value)
                return str(v).strip() if v is not None else ''
            section=''; detail=''; record_kind='阵容码'; n=0
            for cells in s:
                row=cells[0].row; raw_a=str(cells[0].value or '').strip()
                if luotianyi:
                    if re.match(r'^\d+、',raw_a): section=raw_a; detail=''; record_kind='阵容码'
                    elif raw_a.startswith('上方'):
                        count=3 if '三套' in raw_a else 2
                        for prior in [e for e in entries if e['sourceFile']==name and e['sheet']==s.title and e['originalCategory']==section][-count:]:
                            prior['notes']+='\n\n分节补充：'+raw_a
                    elif cells[0].data_type!='f' and raw_a and not raw_a.isdigit() and row>1:
                        detail=raw_a
                        if s.title=='道馆' and '下方' in raw_a and '配置码' in raw_a: record_kind='配置码'
                        elif s.title=='道馆' and '下方' in raw_a and '队伍码' in raw_a: record_kind='队伍码'
                elif s.max_column>=4 and value(row,1) and value(row,1) not in ['副本','阵容名']:
                    section=value(row,1)
                # 只从原始单元格提取码，不能把合并区域复制成多个引用。
                codes=[(c.column,m.group()) for c in cells if isinstance(c.value,str) and c.data_type!='f' for m in re.finditer(r'\|TA\|[0-9a-fA-F]{32}(?![0-9a-fA-F])',c.value)]
                for col,code in codes:
                    correction=''
                    if luotianyi:
                        title=value(row,2); category={'寮活':'寮活动','普逢':'普通逢魔','极逢':'极逢魔'}.get(s.title,s.title)
                        sub=re.sub(r'^\d+、','',section).split('（')[0].strip()
                        sub={'业原火痴':'业原火','僵尸寮':'僵尸寮通用','首领退治铁鼠':'首领退治'}.get(sub,sub)
                        dungeon=sub
                        if s.title=='普逢':
                            sub=dungeon=title.split('|')[-1]; correction='按工作表“普逢”及B列BOSS修正A2误写的“结界突破”标题'
                        if s.title=='契灵' and title.startswith('薙魂') and sub=='狐火':
                            sub=dungeon='薙魂'; correction='按B列薙魂标题修正A44重复的“狐火”分节'
                        if s.title=='极逢':
                            sub=dungeon=next((x for x in ['蜃气楼','土蜘蛛','荒骷髅','地震鲶'] if x in section),'通用')
                        if sub=='英杰试炼':
                            actor='源赖光' if '源赖光' in detail else '藤原道长'
                            dungeon=actor+' · '+('技能本' if '技能' in detail else '经验本')
                        paths=[route(category,sub,dungeon)]
                        notes='\n\n'.join(x for x in [detail,value(row,5),'面板要求：'+value(row,6) if value(row,6) else '', '御魂命名：'+value(row,7) if value(row,7) else ''] if x)
                        author=value(row,9); url=value(row,8)
                    else:
                        title=value(row,col-1); group='近期更新' if s.title=='更新内容' else section
                        paths=dashen_paths(group,title); notes=value(row,col+1); author=value(row,col+2); url=value(row,col+3)
                    entries.append({'id':f'excel-{len(sources)}-{s.title}-{row}-{col}','title':title or '未命名阵容','originalCategory':section if luotianyi else group,'classificationPaths':paths,'classificationCorrection':correction,'recordKind':record_kind,'code':code,'notes':notes,'author':author,'sourceUrl':url,'sheet':s.title,'row':row,'cell':s.cell(row,col).coordinate,'sourceFile':name,'sourceKind':'excel'})
                    n+=1
            info={'name':s.title,'rows':s.max_row,'codeOccurrences':n}
            source['sheets'].append(info); sheets.append({**info,'sourceFile':name})
        source.update(occurrences=len(entries)-start,uniqueCodes=len({x['code'] for x in entries[start:]})); sources.append(source); w.close()
    digest=hashlib.sha256(json.dumps([(x['source'],x['sha256']) for x in sources],ensure_ascii=False,separators=(',',':')).encode()).hexdigest()
    out={'source':'两份用户工作簿','sources':sources,'sha256':digest,'sheets':sheets,'occurrences':len(entries),'uniqueCodes':len(set(x['code'] for x in entries)),'entries':entries}
    save(DATA/'excel.json',out)
    print(json.dumps({k:v for k,v in out.items() if k!='entries'},ensure_ascii=False))

def roster():
    all_rows={}; pages=[]; errors=[]; previous=set(); expected=None
    for page in range(1,31):
        m,raw=request(f'{CATALOG}?rarity=0&page={page}&per_page=100&callback=cb',f'roster-page-{page}')
        pages.append(m)
        if m['status']!='received': errors.append(f'page {page}: {m["status"]}'); break
        d=jsonp(raw)
        save(DATA/f'catalog-page-{page}.json',d)
        if d.get('success') is not True: errors.append('success != true'); break
        rows=d.get('data',{})
        if not isinstance(rows,dict): errors.append('unknown data structure'); break
        ids=set(rows)
        if not ids or not ids-previous: errors.append('empty/repeated/no-new page'); break
        all_rows.update(rows); previous.update(ids)
        expected=d.get('total_num',expected)
        total_page=d.get('total_page')
        print('page',page,'items',len(ids),'metadata', {k:v for k,v in d.items() if k!='data'})
        if total_page is not None and page>=int(total_page): break
    else: errors.append('safety-page-limit')
    if expected is not None and len(all_rows)!=int(expected): errors.append(f'count mismatch {len(all_rows)} != {expected}')
    m,raw=request(STATIC,'static-roster')
    static=jsonp(raw) if m['status']=='received' else []
    save(DATA/'static-roster.json',static)
    names={str(x.get('id')):x for x in static} if isinstance(static,list) else static
    union=set(all_rows)|set(names)
    result=[]
    rarity_map={1:'N',2:'R',3:'SR',4:'SSR',5:'SP',6:'UR'}
    for i in sorted(union,key=lambda a:int(a) if a.isdigit() else 0):
        a=all_rows.get(i,{}); b=names.get(i,{})
        rarity=a.get('rarity')
        result.append({'id':i,'name':a.get('name',b.get('name',f'未知 ID {i}')),'rarityRaw':rarity,'rarity':rarity_map.get(rarity,'未知'),'isCollaboration':a.get('interactive')==1,'isGua':a.get('material_type')==101,'isMaterial':bool(a.get('material_type')) and a.get('material_type')!=101,'official':a,'static':b,'awakeningAvailability':'unknown','releasedAt':None,'releasedBeforeCutoffVerified':False,'releaseEvidenceUrls':[],'variants':[]})
    report={'collectedAt':now(),'cutoffDate':'2026-09-12','region':'CN_MAINLAND_LIVE','rosterCompleteness':'unverified','dynamicCount':len(all_rows),'staticCount':len(names),'unionCount':len(union),'onlyDynamic':sorted(set(all_rows)-set(names)),'onlyStatic':sorted(set(names)-set(all_rows)),'paginationErrors':errors,'pages':len(pages),'latestSentinel':[x for x in result if x['name']=='石长姬'],'rarities':dict(collections.Counter(x['rarity'] for x in result))}
    save(DATA/'roster.json',result); save(DATA/'roster-report.json',report)
    print(json.dumps(report,ensure_ascii=False))

FAMILIES={'art-before':('shishen_big_beforeAwake','png','unawakened'),'art-after':('shishen_big_afterAwake','png','awakened'),'portrait-before':('before_awake','jpg','unawakened'),'portrait-after':('after_awake','jpg','awakened'),'bookmark':('mark_btn','png','unknown'),'icon':('shishen','png','unknown')}

def image_probe(ids,families,report_name='sample-assets.json'):
    from PIL import Image
    output=[]
    for i in ids:
        if not str(i).isdigit(): raise ValueError('ID must be numeric')
        for family in families:
            folder,ext,state=FAMILIES[family]
            m,raw=request(f'{CDN}/{folder}/{i}.{ext}',f'image-{i}-{family}')
            item={**m,'id':str(i),'family':family,'requestedState':state,'verifiedState':None,'contentKind':'unverified','identityVerified':False,'stateVerified':False,'nativeCardVerified':False}
            if m['status']=='received':
                try:
                    im=Image.open(io.BytesIO(raw)); im.verify()
                    im=Image.open(io.BytesIO(raw)); im.load()
                    if im.width*im.height>24000000: raise ValueError('pixel-limit')
                    item.update(width=im.width,height=im.height,format=im.format,alpha='A' in im.getbands(),status='downloaded_valid_image')
                    dest=DATA/'images'/str(i)/(family+'.'+ext)
                    dest.parent.mkdir(parents=True,exist_ok=True); dest.write_bytes(raw)
                    item['localPath']=dest.relative_to(ROOT).as_posix()
                except Exception as e: item.update(status='invalid_image',error=str(e)[:200])
            output.append(item)
            print(i,family,item['status'],item.get('width'),item.get('height'),flush=True)
            if report_name: save(DATA/report_name,output)
            if item['status'] in ['access_denied','rate_limited','network_error']: return output
    return output

def awakening_status(row, attrs):
    # The official shishen page template suppresses after-awake art for these
    # rarities and ID 401, even when the attribute API returns a dictionary.
    if row.get('rarity') in ['N','SP','UR'] or row['id']=='401':
        return 'not_applicable'
    return 'supported' if isinstance(attrs.get('1'),dict) and attrs['1'] else 'not_applicable' if '1' in attrs and attrs['1'] in [None,''] else 'unknown'

def collect_full():
    rows=read(DATA/'roster.json'); done=read(DATA/'asset-manifest.json',{})
    ordered=sorted(rows,key=lambda x:(x['id'] not in ['608','607','593','200','217','203','205','605'], -int(x['id'])))
    for n,row in enumerate(ordered):
        i=row['id']
        if i in done: continue
        attrs={}; evidence=[]
        for awake in [0,1]:
            url=CATALOG.replace('get_heroid_list','get_hero_attr')+f'?heroid={i}&awake={awake}&level=40&star=6&callback=cb'
            m,raw=request(url,f'attr-{i}-{awake}')
            evidence.append(m)
            if m['status'] in ['access_denied','rate_limited','network_error']: return
            if m['status']=='received':
                d=jsonp(raw)
                if d.get('success') is True: attrs[str(awake)]=d.get('data')
        state=awakening_status(row,attrs)
        families=['art-before','portrait-before']+(['art-after','portrait-after'] if state=='supported' else [])
        variants=image_probe([i],families,None)
        for a in variants:
            valid=a['status']=='downloaded_valid_image'
            a.update(contentKind='official_art' if a['family'].startswith('art') else 'official_portrait', identityVerified=valid, stateVerified=valid and state!='unknown',verifiedState=a['requestedState'] if valid and state!='unknown' else None,visualReview='sample-reviewed' if i in ['608','217'] else 'source-mapped-not-individually-reviewed')
        valid_images=[a for a in variants if a['status']=='downloaded_valid_image']
        duplicated=len(set(a.get('sha256') for a in valid_images))!=len(valid_images)
        if duplicated:
            for a in variants: a.update(stateVerified=False,verifiedState=None,duplicateReviewRequired=True)
        done[i]={'awakeningAvailability':state,'awakeningEvidence':evidence,'baseAttrs40':attrs,'variants':variants,'nativeCardStatus':'missing_unverified','duplicates':duplicated}
        save(DATA/'asset-manifest.json',done)
        print('COMPLETE',n+1,len(rows),row['name'],state,flush=True)
        if any(a['status'] in ['access_denied','rate_limited','network_error'] for a in variants): return

def decode(all_codes=False):
    codes=list(dict.fromkeys(x['code'] for x in read(DATA/'excel.json')['entries']))
    if not all_codes: codes=codes[:1]+['|TA|9f0305306516a5b9f1b84a1e487e3385']
    results=read(DATA/'decoded.json',{})
    for n,code in enumerate(codes):
        key=hashlib.sha256(code.encode()).hexdigest()[:20]
        m,raw=request(API,'decode-'+key,{'teamCode':code},1024*1024)
        try: payload=jsonp(raw)
        except Exception: payload=None
        data=payload.get('data') if isinstance(payload,dict) else None
        valid=isinstance(data,dict) and isinstance(data.get('entities'),list) and isinstance(data.get('editableTargets'),list) and isinstance(data.get('slotCount'),int)
        results[code]={'evidence':m,'payload':payload,'state':'decoded' if valid and payload.get('ok') is True else 'failed','contentCrossChecked':False}
        save(DATA/'decoded.json',results)
        print(n+1,len(codes),results[code]['state'],m['httpStatus'],flush=True)
        if not all_codes: print(json.dumps(payload,ensure_ascii=False)[:15000])
        if m['status'] in ['access_denied','rate_limited','network_error']: break

if __name__=='__main__':
    p=argparse.ArgumentParser(); p.add_argument('task',choices=['excel','roster','sample','decode','decode-all','full']); a=p.parse_args()
    if a.task=='excel': extract_excel()
    if a.task=='roster': roster()
    if a.task=='sample': image_probe([200,217,593,400,205],list(FAMILIES))
    if a.task in ['decode','decode-all']: decode(a.task=='decode-all')
    if a.task=='full': collect_full()
